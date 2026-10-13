import base64
import uuid
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from openpyxl import load_workbook
from app import models as m
from app.core import time as timezone
from app.db.session import get_session, SessionLocal, active_session
from app.services.business import execute, RuleError, debt, cash_balance
from app.services.projection import snapshot
from app.services.finance import excel
from tests_fastapi.support import TestCase


class FinanceTests(TestCase):
    def setUp(self):
        self.branch=m.Branch.objects.create(name='Finance fixture')
        self.admin=m.User.objects.create_user(username='finance-admin@test.com',email='finance-admin@test.com',role='admin',branch=self.branch)
        self.staff=m.User.objects.create_user(username='finance-staff@test.com',email='finance-staff@test.com',role='staff',branch=self.branch)
        plan=m.Plan.objects.create(branch=self.branch,name='Finance package',price=Decimal('100.01'),duration_value=30)
        self.act('addMember',{'name':'Finance member','phone':'05329998877','consent':True,'planId':str(plan.id),'installments':2})
        self.invoice=m.Invoice.objects.get(); self.member=m.Member.objects.get(); self.sub=m.Subscription.objects.get()

    def act(self,action,p,user=None,key=None): return execute(action,p,user or self.admin,key or uuid.uuid4())
    def collect(self,value,**extra): return self.act('collect',{'invoiceId':str(self.invoice.id),'amount':value,'method':'cash',**extra},self.staff)
    def refund(self,payment,value,user=None): return self.act('refund',{'paymentId':str(payment.id),'amount':value,'reason':'Payment return'},user)

    def test_full_partial_multiple_and_exact_amounts(self):
        self.collect('0.01'); self.collect('40.00'); self.assertEqual(debt(self.invoice),6000)
        self.collect('60.00'); self.assertEqual(debt(self.invoice),0)
        with self.assertRaises(RuleError): self.collect('0.01')
        self.assertEqual(sum(r.amount_cents for r in m.Payment.objects.all()),10001)

    def test_partial_full_refund_and_unauthorized_refund(self):
        self.collect('100.01'); payment=m.Payment.objects.get()
        with self.assertRaises(RuleError): self.refund(payment,'1',self.staff)
        self.refund(payment,'20.01'); self.assertEqual(debt(self.invoice),2001)
        self.refund(payment,'80.00'); self.assertEqual(debt(self.invoice),10001)
        with self.assertRaises(RuleError): self.refund(payment,'0.01')
        self.assertEqual(m.Payment.objects.count(),1); self.assertEqual(m.Refund.objects.count(),2)

    def test_idempotent_collection_and_refund(self):
        payload={'invoiceId':str(self.invoice.id),'amount':'10.01','method':'pos'}; key=uuid.uuid4()
        self.act('collect',payload,key=key); self.act('collect',payload,key=key)
        self.assertEqual(m.Payment.objects.count(),1)
        payment=m.Payment.objects.get(); key=uuid.uuid4();payload={'paymentId':str(payment.id),'amount':'5','reason':'Return duplicate'}
        self.act('refund',payload,key=key); self.act('refund',payload,key=key); self.assertEqual(m.Refund.objects.count(),1)

    def test_installment_revision_preserves_old_rows_and_marks_overdue(self):
        self.collect('10.01')
        yesterday=(timezone.localdate()-timedelta(days=1)).isoformat(); later=(timezone.localdate()+timedelta(days=10)).isoformat()
        self.act('saveInstallments',{'invoiceId':str(self.invoice.id),'rows':[{'due':yesterday,'amount':'50.01'},{'due':later,'amount':'50.00'}],'reason':'New agreement'})
        self.assertEqual(m.Installment.objects.filter(archived=True).count(),2)
        self.assertEqual(m.Installment.objects.filter(archived=False).count(),2)
        data=snapshot(self.admin)['finance']; self.assertEqual(data['overdueCents'],4000);self.assertEqual(data['remainingCents'],9000);self.assertEqual(data['pendingInstallments'],2)
        with self.assertRaises(RuleError): self.act('saveInstallments',{'invoiceId':str(self.invoice.id),'rows':[{'due':later,'amount':'1'}],'reason':'Wrong sum'})
        self.assertEqual(m.Installment.objects.filter(archived=False).count(),2)

    def test_specific_invoice_and_snapshot_persistence(self):
        plan=m.Plan.objects.get(); self.act('renew',{'memberId':str(self.member.id),'planId':str(plan.id)})
        other=m.Invoice.objects.exclude(pk=self.invoice.id).first()
        self.act('collect',{'invoiceId':str(other.id),'amount':'25','method':'bank'})
        self.assertEqual(debt(self.invoice),10001);self.assertEqual(debt(other),7501)
        self.client.force_login(self.admin)
        first=self.client.get('/api/v1/state/').json()['db'];second=self.client.get('/api/v1/state/').json()['db']
        self.assertEqual(first['payments'],second['payments']);self.assertEqual(first['finance']['remainingCents'],17502)

    def test_backdated_payment_cash_uses_recording_time(self):
        self.invoice.created_at=timezone.now()-timedelta(days=10);self.invoice.save()
        session=m.CashSession.objects.create(branch=self.branch,actor=self.admin,opening_cents=0,opened_at=timezone.now()-timedelta(seconds=2))
        self.collect('30',at=(timezone.now()-timedelta(days=2)).isoformat()); payment=m.Payment.objects.get()
        self.assertGreater(payment.recorded_at,payment.at)
        self.assertEqual(cash_balance(self.branch,session,timezone.now()),3000)
        self.refund(payment,'10');self.assertEqual(cash_balance(self.branch,session,timezone.now()),2000)

    def test_metadata_correction_and_closed_cash_protection(self):
        self.collect('30'); payment=m.Payment.objects.get()
        payload={'paymentId':str(payment.id),'method':'bank','note':'=unsafe note','at':payment.at.isoformat(),'reason':'Wrong method'}
        with self.assertRaises(RuleError):self.act('correctPayment',payload,self.staff)
        m.CashSession.objects.create(branch=self.branch,actor=self.admin,opened_at=payment.recorded_at-timedelta(seconds=1),closed_at=payment.recorded_at+timedelta(seconds=1),opening_cents=0)
        with self.assertRaises(RuleError):self.act('correctPayment',payload)
        payload['method']='cash';self.act('correctPayment',payload)
        self.assertEqual(m.Payment.objects.get().note,'=unsafe note');self.assertEqual(m.Audit.objects.filter(type='correctPayment').count(),1)
        with self.assertRaises(RuleError):self.act('correctPayment',{**payload,'amount':'10'})

    def test_cancellation_requires_ended_membership_and_refunded_sale(self):
        payload={'invoiceId':str(self.invoice.id),'reason':'Sale cancelled','confirmed':True}
        with self.assertRaises(RuleError):self.act('cancelInvoice',payload)
        self.sub.status='cancelled';self.sub.save();self.collect('10')
        with self.assertRaises(RuleError):self.act('cancelInvoice',payload)
        self.refund(m.Payment.objects.get(),'10');self.act('cancelInvoice',payload)
        self.assertEqual(debt(self.invoice),0);self.assertEqual(m.Payment.objects.count(),1);self.assertEqual(m.Refund.objects.count(),1)

    def test_excel_filters_decimal_and_formula_safety(self):
        self.collect('10.01',method='pos',note='=1+2')
        result=excel(self.admin,{'method':'pos'});book=load_workbook(BytesIO(base64.b64decode(result['base64'])))
        sheet=book.active;self.assertEqual(sheet.max_row,2);self.assertEqual(sheet['F2'].value,10.01);self.assertEqual(sheet['I2'].data_type,'s')
        self.assertEqual(load_workbook(BytesIO(base64.b64decode(excel(self.admin,{'method':'bank'})['base64']))).active.max_row,1)
        self.client.force_login(self.staff);self.assertEqual(self.client.get('/api/v1/finance/export/').status_code,403)

    def test_concurrent_collection_cannot_overpay(self):
        get_session().commit(); gate=Barrier(2);user_id=self.staff.id;invoice_id=str(self.invoice.id)
        def attempt():
            with SessionLocal() as session:
                token=active_session.set(session)
                try:
                    user=session.get(m.User,user_id);gate.wait(timeout=10)
                    execute('collect',{'invoiceId':invoice_id,'amount':'80','method':'cash'},user,uuid.uuid4());session.commit();return 'ok'
                except RuleError:session.rollback();return 'rejected'
                finally:active_session.reset(token)
        with ThreadPoolExecutor(max_workers=2) as pool: outcomes=list(pool.map(lambda _:attempt(),range(2)))
        self.assertEqual(sorted(outcomes),['ok','rejected']);get_session().expire_all();self.assertEqual(debt(self.invoice),2001)

    def test_invalid_dates_money_and_finance_permissions(self):
        for value in ['1.001','NaN','-1','100.02']:
            with self.assertRaises(RuleError):self.collect(value)
        with self.assertRaises(RuleError):self.collect('1',at=(timezone.now()+timedelta(days=1)).isoformat())
        with self.assertRaises(RuleError):self.act('saveInstallments',{'invoiceId':str(self.invoice.id),'rows':[],'reason':'Forbidden'},self.staff)
