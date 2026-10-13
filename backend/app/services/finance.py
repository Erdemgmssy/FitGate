"""Finance extensions on the existing invoices/payments/refunds; money stays exact cents."""
import base64
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from app import models as m
from app.core import time as timezone
from .business import require, get, amount, parse_time, parse_day, net_paid
from .projection import wire


def payment_date(p, invoice):
    now=timezone.now()
    at=parse_time(p['at']) if p.get('at') else now
    require(at<=now+timedelta(seconds=30),'Ödeme tarihi gelecekte olamaz.')
    require(timezone.localdate(at)>=timezone.localdate(invoice.created_at),'Ödeme tarihi satıştan önce olamaz.')
    return at


def apply(action,p,user,branch):
    require(user and user.role=='admin','Finans düzeltmeleri için admin yetkisi gerekli.')
    require(isinstance(p.get('reason'),str) and 3<=len(p['reason'].strip())<=1500,'En az üç karakterlik işlem gerekçesi gerekli.')
    if action=='correctPayment':
        obj=get(m.Payment,p.get('paymentId'),branch); before=wire(obj)
        require(not any(key in p for key in ['amount','invoiceId','actorId']),'Tutar/satış/personel değiştirilemez; gerekçeli iade ve yeni tahsilat kullanın.')
        require(p.get('method') in ['cash','bank','pos','mixed'],'Ödeme yöntemi geçersiz.')
        # Method changes must not rewrite an already closed cash ledger or refund route.
        if obj.method!=p['method']:
            require(not m.Refund.objects.filter(payment=obj).exists(),'İade edilmiş ödemenin yöntemi değiştirilemez.')
            if 'cash' in [obj.method,p['method']]:
                require(not m.CashSession.objects.filter(branch=branch,closed_at__isnull=False,opened_at__lte=obj.recorded_at,closed_at__gte=obj.recorded_at).exists(),'Kapanmış kasa döneminin ödeme yöntemi değiştirilemez.')
        obj.method=p['method']; obj.note=str(p.get('note',obj.note))[:600]
        obj.at=payment_date({**p,'at':p.get('at') or obj.at.isoformat()},obj.invoice); obj.save()
    else:
        obj=get(m.Invoice,p.get('invoiceId'),branch); before=wire(obj)
        if action=='cancelInvoice':
            require(p.get('confirmed') is True,'Satış iptalini açıkça onaylayın.')
            require(obj.status!='cancelled' and net_paid(obj)==0,'Satış iptali için önce tüm net tahsilat iade edilmeli.')
            require(obj.subscription.status=='cancelled' if obj.subscription_id else m.Booking.objects.filter(pk=obj.booking_id,status__in=['cancelled','excused']).exists(),'Önce ilgili üyeliği veya ders rezervasyonunu sonlandırın.')
            obj.status='cancelled'; obj.save()
        else:
            require(obj.status!='cancelled','İptal satışa taksit planı oluşturulamaz.')
            rows=p.get('rows')
            require(isinstance(rows,list) and 1<=len(rows)<=36,'Taksit planı 1-36 satır içermeli.')
            parsed=[]
            for row in rows:
                require(isinstance(row,dict),'Taksit satırı geçersiz.')
                due=parse_day(row.get('due')); cents=amount(row.get('amount'))
                require(cents>0,'Taksit tutarı sıfırdan büyük olmalı.')
                parsed.append((due,cents))
            require(sum(value for _,value in parsed)==obj.total_cents,'Taksit toplamı satış tutarına eşit olmalı; tahsilatlar vade sırasıyla uygulanır.')
            require(len({day for day,_ in parsed})==len(parsed),'Taksit vadeleri birbirinden farklı olmalı.')
            before['installments']=[wire(row) for row in m.Installment.objects.filter(invoice=obj,archived=False)]
            m.Installment.objects.filter(invoice=obj,archived=False).update(archived=True)
            for due,cents in sorted(parsed): m.Installment.objects.create(branch=branch,invoice=obj,due=due,amount_cents=cents)
    m.Audit.objects.create(branch=branch,actor=user,type=action,entity_id=str(obj.id),detail=p,before=before,after=wire(obj))
    return {'message':'Finans işlemi kaydedildi.'}


def summary(db):
    today=timezone.localdate().isoformat(); month=today[:7]
    payment_totals={}; refunds={}
    for row in db['refunds']: refunds[row['paymentId']]=refunds.get(row['paymentId'],0)+row['amountCents']
    for row in db['payments']:
        payment_totals[row['invoiceId']]=payment_totals.get(row['invoiceId'],0)+row['amountCents']-refunds.get(row['id'],0)
    remaining=overdue=pending=0
    for invoice in db['invoices']:
        if invoice['status']=='cancelled': continue
        paid=payment_totals.get(invoice['id'],0)
        remaining+=max(0,invoice['totalCents']-paid)
        for part in sorted((r for r in db['installments'] if r['invoiceId']==invoice['id'] and not r.get('archived')),key=lambda r:(r['due'],r['id'])):
            allocated=min(part['amountCents'],max(0,paid)); paid-=allocated
            balance=part['amountCents']-allocated
            if balance>0: pending+=1
            if part['due']<today: overdue+=balance
    def day(row): return timezone.localdate(parse_time(row['at'])).isoformat()
    product_cash=sum(r['totalCents'] for r in db.get('sales',[]) if r['method']=='cash')-sum(r.get('refundCents',0) for r in db.get('stockMovements',[]) if r.get('method')=='cash')
    from .stock import summary as product_summary
    product=product_summary(db)
    return {'productTodayNetCents':product['todayNetCents'],'productMonthNetCents':product['monthNetCents'],'membershipMonthNetCents':sum(r['amountCents'] for r in db['payments'] if day(r).startswith(month) and any(i['id']==r['invoiceId'] and i['subscriptionId'] for i in db['invoices']))-sum(r['amountCents'] for r in db['refunds'] if day(r).startswith(month) and any(p['id']==r['paymentId'] and any(i['id']==p['invoiceId'] and i['subscriptionId'] for i in db['invoices']) for p in db['payments'])),
            'todayCents':sum(r['amountCents'] for r in db['payments'] if day(r)==today),
            'monthCents':sum(r['amountCents'] for r in db['payments'] if day(r).startswith(month)),
            'remainingCents':remaining,'overdueCents':overdue,'pendingInstallments':pending,
            'refundedCents':sum(r['amountCents'] for r in db['refunds']),
            'cashCents':product_cash+sum(r['amountCents'] for r in db['payments'] if r['method']=='cash')-sum(r['amountCents'] for r in db['refunds'] if any(p['id']==r['paymentId'] and p['method']=='cash' for p in db['payments']))}


def excel(user, filters):
    require(user.role=='admin','Rapor dışa aktarımı için admin yetkisi gerekli.')
    from openpyxl import Workbook
    first=parse_day(filters['from']) if filters.get('from') else None
    last=parse_day(filters['to']) if filters.get('to') else None
    require(not first or not last or first<=last,'Tarih aralığı geçersiz.')
    workbook=Workbook(); sheet=workbook.active; sheet.title='Tahsilatlar'
    sheet.append(['İşlem no','Satış no','Üye','Tarih','Yöntem','Tutar TL','İade TL','Personel','Açıklama'])
    for row in m.Payment.objects.filter(branch_id=user.branch_id).order_by('at','pk'):
        day=timezone.localdate(row.at)
        if first and day<first or last and day>last: continue
        if filters.get('memberId') and str(row.invoice.member_id)!=filters['memberId']: continue
        if filters.get('actorId') and str(row.actor_id)!=filters['actorId']: continue
        if filters.get('method') and row.method!=filters['method']: continue
        if filters.get('onlyDebt')=='true':
            from .business import debt
            if debt(row.invoice)<=0: continue
        refund=sum(r.amount_cents for r in m.Refund.objects.filter(payment=row))
        values=[row.number,row.invoice.number,row.invoice.member.name,timezone.localtime(row.at).isoformat(),row.method,Decimal(row.amount_cents)/100,Decimal(refund)/100,row.actor.full_name,row.note]
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str): cell.data_type='s'  # Untrusted text must never become an Excel formula.
        for index in [6,7]: sheet.cell(sheet.max_row,index).number_format='#,##0.00'
    sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
    for column in ['A','B','C','D','E','F','G','H','I']: sheet.column_dimensions[column].width=24
    output=BytesIO(); workbook.save(output)
    return {'filename':'FitGate-tahsilatlar.xlsx','base64':base64.b64encode(output.getvalue()).decode('ascii')}
